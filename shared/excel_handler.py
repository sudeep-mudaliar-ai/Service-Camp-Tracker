"""Excel persistence layer for Service Camp Tracker.

Stores asset entries in an Excel workbook that lives inside a locally
synced OneDrive folder. Uses pandas for reading and openpyxl for
appending rows, with retry handling for transient file locks (e.g.
when OneDrive is syncing or the file is open in Excel).
"""

from __future__ import annotations

import logging
import time
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from shared import config

logger = logging.getLogger(__name__)

# How many times to retry when the workbook is locked, and the pause
# between attempts (seconds).
_LOCK_RETRY_ATTEMPTS: int = 5
_LOCK_RETRY_DELAY_SECONDS: float = 1.0


class ExcelLockedError(Exception):
    """Raised when the workbook stays locked after all retry attempts."""


class ExcelHandler:
    """Encapsulates all reads/writes against the AssetEntries workbook."""

    def __init__(
        self,
        file_path: Path | str = config.ONEDRIVE_EXCEL_PATH,
        sheet_name: str = config.EXCEL_SHEET_NAME,
        columns: list[str] | None = None,
    ) -> None:
        self.file_path = Path(file_path)
        self.sheet_name = sheet_name
        self.columns = columns or list(config.EXCEL_COLUMNS)

    # ------------------------------------------------------------------
    # Workbook lifecycle
    # ------------------------------------------------------------------
    def ensure_file_exists(self) -> None:
        """Create the workbook (with header row) if it does not exist yet."""
        if self.file_path.exists():
            return

        logger.info("Excel file not found, creating: %s", self.file_path)
        self.file_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self.sheet_name
        sheet.append(self.columns)
        self._save_with_retry(workbook)
        logger.info("Created workbook with sheet '%s'", self.sheet_name)

    # ------------------------------------------------------------------
    # Reads
    # ------------------------------------------------------------------
    def read_entries(self) -> pd.DataFrame:
        """Return all entries as a DataFrame (empty frame if file missing)."""
        if not self.file_path.exists():
            return pd.DataFrame(columns=self.columns)

        for attempt in range(1, _LOCK_RETRY_ATTEMPTS + 1):
            try:
                frame = pd.read_excel(
                    self.file_path,
                    sheet_name=self.sheet_name,
                    dtype=str,
                    engine="openpyxl",
                )
                # Normalise: guarantee all expected columns exist
                for column in self.columns:
                    if column not in frame.columns:
                        frame[column] = ""
                return frame[self.columns].fillna("")
            except PermissionError:
                logger.warning(
                    "Excel file locked while reading (attempt %d/%d)",
                    attempt,
                    _LOCK_RETRY_ATTEMPTS,
                )
                time.sleep(_LOCK_RETRY_DELAY_SECONDS)

        raise ExcelLockedError(
            f"Could not read '{self.file_path}': file is locked by another "
            f"process (is it open in Excel?)."
        )

    def barcode_exists(self, barcode: str) -> bool:
        """Return True if the barcode already has a row in the sheet."""
        self.ensure_file_exists()
        frame = self.read_entries()
        if frame.empty:
            return False
        existing = frame["Barcode"].astype(str).str.strip()
        return barcode.strip() in existing.values

    # ------------------------------------------------------------------
    # Writes
    # ------------------------------------------------------------------
    def append_entry(self, entry: dict[str, str]) -> None:
        """Append one entry row to the sheet.

        `entry` keys must match the configured column names. Missing
        keys are written as empty strings.
        """
        self.ensure_file_exists()

        workbook = self._load_with_retry()
        try:
            if self.sheet_name in workbook.sheetnames:
                sheet = workbook[self.sheet_name]
            else:
                sheet = workbook.create_sheet(self.sheet_name)
                sheet.append(self.columns)

            row = [str(entry.get(column, "")).strip() for column in self.columns]
            sheet.append(row)
            self._save_with_retry(workbook)
            logger.info("Appended entry for barcode %s", entry.get("Barcode"))
        finally:
            workbook.close()

    # ------------------------------------------------------------------
    # Lock-aware helpers
    # ------------------------------------------------------------------
    def _load_with_retry(self) -> Workbook:
        """Open the workbook, retrying briefly if it is locked."""
        last_error: Exception | None = None
        for attempt in range(1, _LOCK_RETRY_ATTEMPTS + 1):
            try:
                return load_workbook(self.file_path)
            except PermissionError as error:
                last_error = error
                logger.warning(
                    "Excel file locked while opening (attempt %d/%d)",
                    attempt,
                    _LOCK_RETRY_ATTEMPTS,
                )
                time.sleep(_LOCK_RETRY_DELAY_SECONDS)
        raise ExcelLockedError(
            f"Could not open '{self.file_path}': file is locked. "
            f"Close it in Excel and try again."
        ) from last_error

    def _save_with_retry(self, workbook: Workbook) -> None:
        """Save the workbook, retrying briefly if it is locked."""
        last_error: Exception | None = None
        for attempt in range(1, _LOCK_RETRY_ATTEMPTS + 1):
            try:
                workbook.save(self.file_path)
                return
            except PermissionError as error:
                last_error = error
                logger.warning(
                    "Excel file locked while saving (attempt %d/%d)",
                    attempt,
                    _LOCK_RETRY_ATTEMPTS,
                )
                time.sleep(_LOCK_RETRY_DELAY_SECONDS)
        raise ExcelLockedError(
            f"Could not save '{self.file_path}': file is locked. "
            f"Close it in Excel and try again."
        ) from last_error
